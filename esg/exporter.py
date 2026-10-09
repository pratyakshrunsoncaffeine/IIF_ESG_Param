"""CSV, JSON and Excel output with summary and error evidence."""
from __future__ import annotations

import csv
import json
from collections import defaultdict
from pathlib import Path

COLUMNS = [
    ("Ticker", "ticker"), ("Company", "company"), ("Document Type", "document_type"),
    ("Fiscal Year", "fiscal_year"), ("Source File", "source_file"),
    ("Page Number", "page_number"),
    ("ESG Pillar", "pillar"), ("ESG Topic", "topic"),
    ("Matched Keywords", "matched_keywords"), ("Primary Keyword", "keyword"),
    ("Claim Type", "claim_type"), ("Extracted Statement", "matched_sentence"),
    ("Context Before", "context_before"), ("Context After", "context_after"),
    ("Value", "value"), ("Unit", "unit"), ("Reporting Period", "reporting_period"),
    ("Previous Value", "previous_value"), ("Previous Period", "previous_period"),
    ("Current Value", "current_value"), ("Current Period", "current_period"),
    ("Target Metric", "target_metric"), ("Target Value", "target_value"),
    ("Target Year", "target_year"), ("Baseline Year", "baseline_year"),
    ("Direction", "direction"), ("Confidence", "confidence"),
    ("OCR Used", "ocr_used"), ("Table Index", "table_index"),
    ("Table Row", "table_row"), ("Column Headers", "column_headers"),
    ("Raw Table Data", "raw_table_data"),
]
ERROR_COLUMNS = ["ticker", "source_file", "page_number", "error_type", "message"]


def _cell(record: dict, key: str):
    value = record.get(key)
    return "; ".join(value) if isinstance(value, list) else value


def _sheet(workbook, title: str, headers: list[str], rows: list[list]):
    from openpyxl.styles import Font, PatternFill

    ws = workbook.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24466B")
    for column in ws.columns:
        letter = column[0].column_letter
        width = min(65, max(14, max(len(str(c.value or "")) for c in list(column)[:101]) + 2))
        ws.column_dimensions[letter].width = width
    return ws


def export_results(records: list[dict], errors: list[dict], output_dir: str | Path) -> dict:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / "esg_results.csv"
    json_path = output_dir / "esg_results.json"
    xlsx_path = output_dir / "esg_results.xlsx"
    error_path = output_dir / "extraction_errors.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([heading for heading, _ in COLUMNS])
        writer.writerows([_cell(r, key) for _, key in COLUMNS] for r in records)
    json_path.write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
    with error_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=ERROR_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(errors)
    try:
        from openpyxl import Workbook
    except ImportError:
        # Keep CSV/JSON exports and the app usable if Streamlit is still
        # resolving optional workbook dependencies during a deployment.
        return {"csv": csv_path, "json": json_path, "errors": error_path}

    wb = Workbook()
    wb.remove(wb.active)
    headers = [h for h, _ in COLUMNS]
    to_rows = lambda items: [[_cell(r, key) for _, key in COLUMNS] for r in items]
    _sheet(wb, "All_ESG_Data", headers, to_rows(records))
    for pillar in ("Environmental", "Social", "Governance"):
        _sheet(wb, pillar, headers, to_rows([r for r in records if r["pillar"] == pillar]))
    summary = defaultdict(lambda: {"count": 0, "measured": 0, "targets": 0, "adverse": 0, "pages": set()})
    for r in records:
        entry = summary[(r["company"], r["pillar"], r["topic"])]
        entry["count"] += 1
        entry["measured"] += r["claim_type"] == "MEASURED_RESULT"
        entry["targets"] += r["claim_type"] == "TARGET"
        entry["adverse"] += r["claim_type"] == "ADVERSE_EVENT"
        entry["pages"].add(f'{r["source_file"]}:p{r["page_number"]}')
    summary_rows = [[*key, v["count"], v["measured"], v["targets"], v["adverse"],
                     "; ".join(sorted(v["pages"]))] for key, v in sorted(summary.items())]
    _sheet(wb, "ESG_Summary", ["Company", "Pillar", "Topic", "Number_of_Matches",
                               "Number_of_Measured_Results", "Number_of_Targets",
                               "Number_of_Adverse_Events", "Pages_Found"], summary_rows)
    _sheet(wb, "Extraction_Errors", ERROR_COLUMNS,
           [[e.get(key) for key in ERROR_COLUMNS] for e in errors])
    wb.save(xlsx_path)
    return {"csv": csv_path, "json": json_path, "xlsx": xlsx_path, "errors": error_path}
