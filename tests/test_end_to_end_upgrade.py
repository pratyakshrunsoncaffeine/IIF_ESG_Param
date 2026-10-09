"""End-to-end acceptance checks for grouped reports, article screening and storage."""
from __future__ import annotations

import csv
import json
import sys
from types import SimpleNamespace

import pandas as pd
import pytest

import esg_engine
from config import DEFAULTS
from esg.database import read_table, save_evidence_batch, save_score_batch
from esg.documents import group_documents
from esg.exporter import export_results
from esg.keyword_matcher import KeywordMatcher
from esg.pipeline import process_pdf_bytes
from esg.streamlit_adapter import records_to_score_evidence
from tools.acceptance_runner import run_app_acceptance


def _pdf(text: str) -> bytes:
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text, fontsize=10)
    content = document.tobytes()
    document.close()
    return content


def test_app_acceptance_runner_covers_annual_only_and_grouped_brsr_inputs():
    result = run_app_acceptance()
    assert result["annual_only_tickers"] == ["HDFCBANK.NS"]
    assert result["grouped_document_types"] == ["annual_report", "brsr"]
    assert result["news_calls_per_company"] == ["HDFCBANK.NS", "ICICIBANK.NS"]


def test_grouped_annual_and_optional_brsr_keep_provenance_in_scores_exports_and_database(tmp_path):
    climate = "Scope 1 emissions fell 12% in FY2025 after an audited baseline review."
    social = "Employee safety training reached 94% of staff in FY2025 and improved coverage."
    annual_hdfc = _pdf(f"HDFC Bank Annual Report FY2025. {climate} {social}")
    annual_icici = _pdf("ICICI Bank Annual Report FY2025. Board oversight included 7 directors in FY2025.")
    brsr_hdfc = _pdf(f"HDFC Bank BRSR FY2025. {climate} Community investment reached 20 million rupees in FY2025.")
    groups = group_documents(
        [
            {"file_name": "hdfc_annual.pdf", "ticker": "HDFCBANK.NS", "company_name": "HDFC Bank",
             "fiscal_year": "FY2025", "content": annual_hdfc},
            {"file_name": "icici_annual.pdf", "ticker": "ICICIBANK.NS", "company_name": "ICICI Bank",
             "fiscal_year": "FY2025", "content": annual_icici},
        ],
        [{"file_name": "hdfc_brsr.pdf", "ticker": "HDFCBANK.NS", "company_name": "HDFC Bank",
          "content": brsr_hdfc}],
    )
    by_ticker = {row["ticker"]: row for row in groups}
    assert len(groups) == 2
    assert [doc["document_type"] for doc in by_ticker["HDFCBANK.NS"]["documents"]] == ["annual_report", "brsr"]
    assert by_ticker["HDFCBANK.NS"]["documents"][1]["fiscal_year"] == "FY2025"

    matcher = KeywordMatcher(esg_engine.ESG_DICTIONARY)
    all_records, all_score_evidence = [], []
    page_rows, errors = [], []
    for company in groups:
        for document in company["documents"]:
            document_errors, page_status = [], []
            records = process_pdf_bytes(
                document["content"], document["source_file"], company["company_name"], matcher,
                DEFAULTS, ocr=False, errors=document_errors, page_status=page_status,
            )
            assert page_status and not document_errors
            assert records, document["source_file"]
            for record in records:
                record.update(ticker=company["ticker"], document_type=document["document_type"],
                              fiscal_year=document["fiscal_year"])
            all_records.extend(records)
            all_score_evidence.extend(records_to_score_evidence(
                records, company["ticker"], company["company_name"], document["source_file"],
                document_type=document["document_type"], fiscal_year=document["fiscal_year"],
            ))
            page_rows.extend({"ticker": company["ticker"], **row} for row in page_status)
            errors.extend(document_errors)

    hdfc_topic = [row for row in all_records if row["ticker"] == "HDFCBANK.NS"
                  and row["topic"] == "Climate Change & GHG Emissions"]
    assert {row["source_file"] for row in hdfc_topic} == {"hdfc_annual.pdf", "hdfc_brsr.pdf"}
    assert {row["document_type"] for row in hdfc_topic} == {"annual_report", "brsr"}
    score_rows = [row for row in all_score_evidence if row["ticker"] == "HDFCBANK.NS"
                  and row["topic"] == "Climate Change & GHG Emissions"]
    assert {(row["report_name"], row["document_type"], row["fiscal_year"]) for row in score_rows} == {
        ("hdfc_annual.pdf", "annual_report", "FY2025"),
        ("hdfc_brsr.pdf", "brsr", "FY2025"),
    }

    companies = [{key: row[key] for key in ("ticker", "company_name", "report_name")} for row in groups]
    company_scores = esg_engine.calculate_document_scores(pd.DataFrame(all_score_evidence), companies)
    assert set(company_scores["ticker"]) == {"HDFCBANK.NS", "ICICIBANK.NS"}
    assert len(company_scores) == 6  # One E/S/G scorecard per company, regardless of document count.
    without_duplicate = [row for row in all_score_evidence if not (
        row["ticker"] == "HDFCBANK.NS" and row["document_type"] == "brsr"
        and row["topic"] == "Climate Change & GHG Emissions")]
    one_source_scores = esg_engine.calculate_document_scores(pd.DataFrame(without_duplicate), companies)
    hdfc_e = company_scores.set_index(["ticker", "pillar"]).loc[("HDFCBANK.NS", "E"), "document_evidence_score_0_7"]
    one_source_e = one_source_scores.set_index(["ticker", "pillar"]).loc[("HDFCBANK.NS", "E"), "document_evidence_score_0_7"]
    assert hdfc_e == one_source_e  # The same topic in two PDFs retains provenance but scores once.

    files = export_results(all_records, errors, tmp_path / "exports")
    with files["csv"].open(encoding="utf-8-sig", newline="") as handle:
        exported_rows = list(csv.DictReader(handle))
    exported_topic = [row for row in exported_rows if row["Ticker"] == "HDFCBANK.NS"
                      and row["ESG Topic"] == "Climate Change & GHG Emissions"]
    assert {(row["Document Type"], row["Fiscal Year"], row["Source File"]) for row in exported_topic} == {
        ("annual_report", "FY2025", "hdfc_annual.pdf"),
        ("brsr", "FY2025", "hdfc_brsr.pdf"),
    }
    assert files["xlsx"].exists()
    from openpyxl import load_workbook
    workbook = load_workbook(files["xlsx"], read_only=True, data_only=True)
    sheet = workbook["All_ESG_Data"]
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    doc_type_col, fiscal_year_col, source_col = (header.index(name) for name in
                                                  ("Document Type", "Fiscal Year", "Source File"))
    exported_xlsx = {
        (row[doc_type_col], row[fiscal_year_col], row[source_col])
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[header.index("Ticker")] == "HDFCBANK.NS"
        and row[header.index("ESG Topic")] == "Climate Change & GHG Emissions"
    }
    assert exported_xlsx == {
        ("annual_report", "FY2025", "hdfc_annual.pdf"),
        ("brsr", "FY2025", "hdfc_brsr.pdf"),
    }
    workbook.close()

    db_path = tmp_path / "acceptance.sqlite3"
    run_id = "grouped-document-run"
    news = pd.DataFrame()
    save_evidence_batch(run_id, news, pd.DataFrame(all_score_evidence), all_records, db_path)
    rankings, detail = esg_engine.apply_incidents_and_rank(company_scores, pd.DataFrame(), news)
    save_score_batch(run_id, rankings, detail, path=db_path)
    stored = read_table("esg_dataset", db_path)
    assert len(stored) == len(all_records)
    stored_payloads = [json.loads(value) for value in stored["payload_json"]]
    stored_hdfc = [row for row in stored_payloads if row.get("ticker") == "HDFCBANK.NS"
                   and row.get("topic") == "Climate Change & GHG Emissions"]
    assert {(row["document_type"], row["fiscal_year"], row["source_file"]) for row in stored_hdfc} == {
        ("annual_report", "FY2025", "hdfc_annual.pdf"),
        ("brsr", "FY2025", "hdfc_brsr.pdf"),
    }
    history = read_table("score_history", db_path)
    assert len(history) == 6
    assert set(rankings["ticker"]) == {"HDFCBANK.NS", "ICICIBANK.NS"}


class _WhitespaceTokenizer:
    model_max_length = 512

    def encode(self, text, add_special_tokens=False):
        return str(text).split()

    def decode(self, token_ids, skip_special_tokens=True):
        return " ".join(token_ids)


class _FakeClassifier:
    tokenizer = _WhitespaceTokenizer()

    def __init__(self, fail_on_call=None):
        self.chunks = []
        self.fail_on_call = fail_on_call

    def __call__(self, texts, **kwargs):
        self.chunks.extend(texts)
        if self.fail_on_call and len(self.chunks) == self.fail_on_call:
            raise RuntimeError("synthetic model interruption")
        negative = 0.90 if "TAIL-REGULATORY-FINDING" in texts[0] else 0.10
        neutral = 0.05 if negative > 0.5 else 0.30
        positive = 1 - negative - neutral
        return [[{"label": "negative", "score": negative},
                 {"label": "neutral", "score": neutral},
                 {"label": "positive", "score": positive}]]


def _news_rows():
    market = (
        "HDFC Bank shares climbed after Nuvama reiterated its buy call and raised the price target. "
        "Brokerage analysts discussed trading volume, valuation and the quarterly earnings outlook. "
        "SBI, a separate bank, faced a regulatory fine following a customer fraud investigation. "
    ) * 4
    # The ESG conduct signal intentionally appears well past both the old 5,000-character cap
    # and the first 512-token model window.
    filler = ("The bank's public earnings update reviewed lending activity, services and operations. " * 190)
    tail = (
        "TAIL-REGULATORY-FINDING. The Reserve Bank of India fined HDFC Bank after finding consumer fraud "
        "that harmed customers through mis-selling. The regulatory fine followed a completed investigation."
    )
    return pd.DataFrame([
        {"ticker": "HDFCBANK.NS", "company_name": "HDFC Bank", "aliases": ["HDFC"],
         "pillar": "G", "queried_pillars": "G", "publisher": "Business Standard",
         "url": "https://www.business-standard.com/markets/hdfc-price-target", "canonical_url": "https://www.business-standard.com/markets/hdfc-price-target",
         "title": "Nuvama raises HDFC Bank price target", "published_at": "2026-08-20"},
        {"ticker": "HDFCBANK.NS", "company_name": "HDFC Bank", "aliases": ["HDFC"],
         "pillar": "G", "queried_pillars": "G", "publisher": "Business Standard",
         "url": "https://www.business-standard.com/finance/hdfc-rbi-order", "canonical_url": "https://www.business-standard.com/finance/hdfc-rbi-order",
         "title": "RBI action against HDFC Bank", "published_at": "2026-08-19"},
    ]), market, filler + tail


def _fetcher_for(market, valid, fail=False):
    def fetch(url, _domains):
        if url.endswith("hdfc-price-target"):
            body = market
        else:
            body = valid
        return {"url": url, "article_text": body, "retrieval_state": "retrieved",
                "retrieval_reason": "fixture full body"}
    return fetch


def test_unrelated_market_candidate_cannot_change_coverage_score_or_penalty(monkeypatch):
    monkeypatch.setitem(sys.modules, "langdetect", SimpleNamespace(detect=lambda _text: "en"))
    candidates, market, valid = _news_rows()
    classifier = _FakeClassifier()
    news = esg_engine.classify_news(candidates, classifier, article_fetcher=_fetcher_for(market, valid))
    unrelated = news.loc[news["url"].str.endswith("hdfc-price-target")].iloc[0]
    regulatory = news.loc[news["url"].str.endswith("hdfc-rbi-order")].iloc[0]
    assert unrelated["relevance_decision"] == "rejected"
    assert not unrelated["scoring_eligible"]
    assert regulatory["relevance_decision"] == "accepted"
    assert regulatory["scoring_eligible"]
    assert regulatory["full_article_available"] and regulatory["model_input_source"] == "full article"
    assert regulatory["token_count"] > 512 and regulatory["chunk_count"] > 1
    assert regulatory["chunks_scored"] == regulatory["chunk_count"]
    assert regulatory["full_text_scored"]
    assert any("TAIL-REGULATORY-FINDING" in chunk for chunk in classifier.chunks)

    monkeypatch.setitem(esg_engine.NEWS_PILLAR_TARGETS, "G", 2)
    scores = pd.DataFrame([{
        "ticker": "HDFCBANK.NS", "company_name": "HDFC Bank", "pillar": pillar,
        "document_evidence_score_0_7": 5.0, "report_name": "annual.pdf",
    } for pillar in ("E", "S", "G")])
    incident = pd.DataFrame([{
        "ticker": "HDFCBANK.NS", "pillar": "G", "summary": "Final regulator finding",
        "status": "final regulator order or sanction", "official_record_url": "https://rbi.org.in/orders/final.pdf",
        "severity_points": 1.0, "analyst_confirmed": True, "already_in_metric": False,
    }])
    with_unrelated, detail_with = esg_engine.apply_incidents_and_rank(scores, incident, news)
    without_unrelated, detail_without = esg_engine.apply_incidents_and_rank(
        scores, incident, news.loc[news["relevance_decision"].ne("rejected")].copy(),
    )
    g_with = detail_with.set_index("pillar").loc["G"]
    g_without = detail_without.set_index("pillar").loc["G"]
    assert g_with["news_articles_found"] == g_without["news_articles_found"] == 1
    assert not g_with["news_target_met"] and not g_without["news_target_met"]
    assert g_with["news_risk_score_0_7"] == g_without["news_risk_score_0_7"] == 3.5
    assert g_with["approved_incident_penalty"] == g_without["approved_incident_penalty"] == 1.0
    assert with_unrelated.iloc[0]["overall_score_0_21"] == without_unrelated.iloc[0]["overall_score_0_21"]


def test_retrieval_and_partial_classifier_failures_never_create_scored_news(monkeypatch):
    monkeypatch.setitem(sys.modules, "langdetect", SimpleNamespace(detect=lambda _text: "en"))
    candidates, market, valid = _news_rows()

    def failed_fetch(url, _domains):
        return {"url": url, "article_text": "", "retrieval_state": "failed", "retrieval_reason": "offline"}

    unavailable = esg_engine.classify_news(candidates.iloc[[1]], _FakeClassifier(), article_fetcher=failed_fetch)
    assert unavailable.iloc[0]["retrieval_state"] == "failed"
    assert not unavailable.iloc[0]["scoring_eligible"]
    assert pd.isna(unavailable.iloc[0]["finbert_negative_probability"])
    assert unavailable.iloc[0]["model_input_source"] == "none: no headline fallback"

    broken_model = _FakeClassifier(fail_on_call=2)
    failed_score = esg_engine.classify_news(
        candidates.iloc[[1]], broken_model, article_fetcher=_fetcher_for(market, valid),
    )
    scored = failed_score.iloc[0]
    assert scored["chunk_count"] > 1
    assert 0 < scored["chunks_scored"] < scored["chunk_count"]
    assert not scored["full_text_scored"] and not scored["scoring_eligible"]
    assert pd.isna(scored["finbert_negative_probability"])
    assert "synthetic model interruption" in scored["scoring_error"]

    base_scores = pd.DataFrame([{
        "ticker": "HDFCBANK.NS", "company_name": "HDFC Bank", "pillar": "G",
        "document_evidence_score_0_7": 5.0,
    }])
    ranking, detail = esg_engine.apply_incidents_and_rank(base_scores, pd.DataFrame(), failed_score)
    assert detail.iloc[0]["news_articles_found"] == 0
    assert not detail.iloc[0]["news_target_met"]
    assert detail.iloc[0]["news_risk_score_0_7"] == 3.5
    assert ranking.iloc[0]["news_articles_found"] == 0
